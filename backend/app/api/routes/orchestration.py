import json
from typing import List, Optional
from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from app.core.json_utils import json_loads
from app.core.config import BACKEND_DIR
from app.db.models import Dataset, SimulationRequest, Trip
from app.dmfe.models import DMFEBatch
from app.schemas.orchestration import DatasetResponse
from app.api.deps import SessionDep, CurrentUser

# Provider, Vehicle, OptimizationResult, OptimizationResultResponse and
# AIOrchestrator were imported here for the legacy /optimize path. That path is
# retired (see run_optimization below) and the OptimizationResult table is no
# longer written or read, so the imports are dropped rather than left implying
# the path is still wired up.

router = APIRouter()


@router.get("/datasets", response_model=List[DatasetResponse])
def list_datasets(db: SessionDep, current_user: CurrentUser):
    return db.query(Dataset).order_by(Dataset.created_at.desc()).all()


@router.post("/datasets/upload", response_model=DatasetResponse)
def upload_dataset(
    name: str = Form(...),
    file_type: str = Form(...),
    data_type: str = Form(...),
    description: str = Form(""),
    file: UploadFile = File(None),
    db: SessionDep = None,
    current_user: CurrentUser = None,
):
    import tempfile, os

    # Normalize data_type (Vehicles -> vehicle)
    normalized_data_type = data_type.lower()
    if normalized_data_type == "vehicles":
        normalized_data_type = "vehicle"

    row_count = 0
    file_path = None
    content = None
    rows = []

    if file:
        content = file.file.read()
        ext = (file.filename or "").lower()
        ftype = file_type.lower()

        if ftype == "csv" or ext.endswith(".csv"):
            try:
                import csv
                from io import StringIO
                reader = csv.DictReader(StringIO(content.decode("utf-8")))
                rows = [dict(r) for r in reader]
                row_count = len(rows)
            except UnicodeDecodeError:
                raise HTTPException(400, "CSV file must be UTF-8 text")
        elif ftype == "json" or ext.endswith(".json"):
            try:
                data = json.loads(content)
            except (json.JSONDecodeError, UnicodeDecodeError):
                raise HTTPException(400, "Invalid JSON file")
            row_count = len(data) if isinstance(data, list) else 1
        elif ftype == "xlsx" or ext.endswith(".xlsx"):
            try:
                import openpyxl
                from io import BytesIO
                wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
                ws = wb.active
                iter_rows = ws.iter_rows(values_only=True)
                header = next(iter_rows, None)
                if header is None:
                    wb.close()
                    raise HTTPException(400, "XLSX file has no data rows")
                cols = [str(c).strip() if c is not None else "" for c in header]
                for values in iter_rows:
                    rows.append(dict(zip(cols, values)))
                wb.close()
                row_count = len(rows)
            except HTTPException:
                raise
            except Exception:
                raise HTTPException(400, "Invalid XLSX file")

    if normalized_data_type == "vehicle" and (file_type.lower() == "csv" or (file and file.filename and file.filename.lower().endswith((".csv", ".xlsx")))):
        if not file or not content:
            raise HTTPException(400, "A CSV or XLSX file is required to import vehicles")

    dataset = Dataset(
        name=name,
        file_type=file_type,
        data_type=data_type,
        file_path=None,
        row_count=row_count,
        description=description,
    )
    db.add(dataset)

    dataset_dir = os.path.join(BACKEND_DIR, "datasets")
    os.makedirs(dataset_dir, exist_ok=True)
    
    try:
        if normalized_data_type == "vehicle" and (file_type.lower() == "csv" or (file and file.filename and file.filename.lower().endswith((".csv", ".xlsx")))):
            from app.db.models import Vehicle, Provider

            provider = db.query(Provider).first()
            if not provider:
                provider = Provider(name="Default Provider", provider_type="Fleet")
                db.add(provider)
                db.flush()

            for row in rows:
                vehicle_id = row.get("vehicle_id", "")
                if not vehicle_id:
                    continue

                vehicle = db.query(Vehicle).filter(Vehicle.name == vehicle_id).first()
                if not vehicle:
                    vehicle = Vehicle(
                        provider_id=provider.id,
                        name=vehicle_id,
                        registration_number=vehicle_id,
                        vehicle_type=row.get("vehicle_type", "Car")
                    )
                    db.add(vehicle)

                if "vehicle_type" in row:
                    vehicle.vehicle_type = row["vehicle_type"]
                if "capacity" in row:
                    try:
                        vehicle.capacity = int(row["capacity"])
                    except ValueError:
                        pass
                if "status" in row:
                    vehicle.status = row["status"]
                if "latitude" in row:
                    try:
                        vehicle.current_lat = float(row["latitude"])
                    except ValueError:
                        pass
                if "longitude" in row:
                    try:
                        vehicle.current_lng = float(row["longitude"])
                    except ValueError:
                        pass
                if "fuel_type" in row:
                    vehicle.fuel_type = row["fuel_type"]
        
        # Finally, persist the file
        if file:
            ext = os.path.splitext(file.filename)[1] if file.filename else ""
            fd, file_path = tempfile.mkstemp(suffix=ext, dir=dataset_dir)
            os.close(fd)
            with open(file_path, "wb") as f:
                f.write(content)
            dataset.file_path = file_path

        db.commit()
        db.refresh(dataset)
        
    except Exception as e:
        db.rollback()
        if file_path and os.path.exists(file_path):
            try:
                os.unlink(file_path)
            except OSError:
                pass
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(500, f"Upload failed: {str(e)}")

    return dataset


@router.delete("/datasets/{dataset_id}")
def delete_dataset(dataset_id: int, db: SessionDep, current_user: CurrentUser):
    import os
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(404, "Dataset not found")
        
    file_path = dataset.file_path
    
    db.delete(dataset)
    db.commit()
    
    if file_path and os.path.exists(file_path):
        try:
            os.unlink(file_path)
        except OSError:
            pass
            
    return {"message": "Dataset deleted"}


@router.post("/optimize")
def run_optimization(db: SessionDep, current_user: CurrentUser):
    """
    RETIRED — this endpoint ran the legacy AIOrchestrator (app/engine/optimizer.py),
    a second optimizer independent of A-DMFE.

    It was actively harmful, for two reasons:

      1. It selected EVERY request with status == "Pending" — not just the ones
         the caller had just simulated — and set them to status = "Optimized"
         (app/engine/optimizer.py::_mark_optimized). No other code in this
         codebase reads the value "Optimized", and the A-DMFE pipeline filters
         on status == "Pending", so each call silently and permanently removed
         the entire A-DMFE queue from the engine's reach.

      2. It wrote OptimizationResult rows, while GET /api/orchestration/results
         reads Trip rows. The page therefore displayed one entity type after
         Run and a different one after Refresh.

    The application now runs a single engine. Route optimisation is performed
    by POST /api/dmfe/analyze, whose output this page reads via
    GET /api/orchestration/results.

    The AIOrchestrator class is left in the tree unmodified for reference; it
    simply has no caller.
    """
    raise HTTPException(
        status_code=410,
        detail=(
            "The legacy orchestrator has been retired — it consumed the A-DMFE "
            "pending queue. Use POST /api/dmfe/analyze instead; its results are "
            "served by GET /api/orchestration/results."
        ),
    )


def _serialize_request(r: SimulationRequest) -> dict:
    if not r:
        return {}
    raw_type = (r.request_type or "").lower()
    if raw_type == "ride":
        service_label = "Passenger"
    elif raw_type == "food":
        service_label = "Food"
    elif raw_type == "parcel":
        service_label = "Parcel"
    else:
        service_label = raw_type.capitalize() or "Passenger"

    def _coords(lat, lng):
        # Guard corrupt rows with NULL coordinates (would otherwise raise
        # TypeError on the `:.4f` format and surface as an unhandled 500).
        if lat is None or lng is None:
            return "Unavailable"
        return f"Lat: {lat:.4f}, Lng: {lng:.4f}"

    return {
        "id": r.id,
        "request_type": raw_type,
        "service_label": service_label,
        "pickup_address": r.pickup_address or _coords(r.pickup_lat, r.pickup_lng),
        "drop_address": r.drop_address or _coords(r.drop_lat, r.drop_lng),
        "pickup_lat": r.pickup_lat,
        "pickup_lng": r.pickup_lng,
        "drop_lat": r.drop_lat,
        "drop_lng": r.drop_lng,
        "priority": r.priority or "Medium",
        "demand": r.demand or 1,
        "weight_kg": r.weight_kg or 0.0,
        "max_acceptable_delay_min": r.max_acceptable_delay_min,
        "requested_time": r.request_timestamp.isoformat() if r.request_timestamp else (r.created_at.isoformat() if r.created_at else None),
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


def _trip_to_result(t: Trip, db: SessionDep) -> dict:
    """
    Serialize one A-DMFE Trip into the rich shape the Orchestration page renders.
    Enriches with requests, batch factors, driver/vehicle details, and stop locations.
    """
    request_ids = json_loads(t.request_ids_json, [])
    stop_order = json_loads(t.stop_order_json, [])

    # Query real SimulationRequest records
    reqs = (
        db.query(SimulationRequest).filter(SimulationRequest.id.in_(request_ids)).all()
        if request_ids
        else []
    )
    req_map = {r.id: r for r in reqs}

    # Enrich stops with address and coordinates
    enriched_stops = []
    for s in stop_order:
        rid = s.get("request_id")
        action = s.get("action", "stop")
        req = req_map.get(rid)
        stop_dict = dict(s)
        if req:
            stop_dict["request_type"] = req.request_type
            raw_type = (req.request_type or "").lower()
            stop_dict["service_label"] = (
                "Passenger" if raw_type == "ride"
                else "Food" if raw_type == "food"
                else "Parcel" if raw_type == "parcel"
                else raw_type.capitalize()
            )
            stop_dict["priority"] = req.priority or "Medium"
            if action == "pickup":
                stop_dict["address"] = req.pickup_address or f"Pickup #{rid}"
                stop_dict["lat"] = req.pickup_lat
                stop_dict["lng"] = req.pickup_lng
            else:
                stop_dict["address"] = req.drop_address or f"Drop #{rid}"
                stop_dict["lat"] = req.drop_lat
                stop_dict["lng"] = req.drop_lng
        enriched_stops.append(stop_dict)

    # Batch and feasibility metadata
    batch = t.batch
    if not batch and t.batch_id:
        batch = db.query(DMFEBatch).filter(DMFEBatch.id == t.batch_id).first()
    if not batch and t.trip_code:
        batch = db.query(DMFEBatch).filter(DMFEBatch.batch_code == t.trip_code).first()

    factor_scores = json_loads(batch.factor_scores_json, {}) if batch and batch.factor_scores_json else {}
    factor_details = json_loads(batch.factor_details_json, {}) if batch and batch.factor_details_json else {}
    reasons = json_loads(batch.reason_json, []) if batch and batch.reason_json else []
    compatibility_score = (
        batch.compatibility_score
        if batch and batch.compatibility_score is not None
        else (t.optimization_score or 100.0)
    )

    # Driver details
    driver_info = None
    if t.driver:
        driver_info = {
            "id": t.driver.id,
            "name": t.driver.name,
            "status": t.driver.status,
            "rating": getattr(t.driver, "rating", 4.8),
            "phone": getattr(t.driver, "phone", None),
            "provider_name": t.driver.provider.name if t.driver.provider else "Default Provider",
        }

    # Vehicle details
    vehicle_info = None
    if t.vehicle:
        vehicle_info = {
            "id": t.vehicle.id,
            "name": t.vehicle.name,
            "capacity": t.vehicle.capacity,
            "fuel_type": getattr(t.vehicle, "fuel_type", "Standard"),
            "plate_number": getattr(t.vehicle, "plate_number", "N/A"),
            "cost_per_km": t.vehicle.cost_per_km or 10.0,
        }

    total_demand = sum(r.demand or 1 for r in reqs)

    return {
        "id": t.id,
        "raw_id": t.id,
        "is_batch": False,
        "batch_id": t.trip_code,
        "decision": "Accepted",
        "status": "Accepted",
        "is_shared": t.is_shared,
        "request_count": len(request_ids) if isinstance(request_ids, list) else 0,
        "requests": [_serialize_request(r) for r in reqs],
        "provider_id": t.driver.provider_id if t.driver else None,
        "vehicle_id": t.vehicle_id,
        "driver": driver_info,
        "vehicle": vehicle_info,
        "current_load": total_demand,
        "vehicle_capacity": t.vehicle.capacity if t.vehicle else None,
        "utilization_pct": t.utilization_pct or 0.0,
        "driver_availability": "Available (Assigned)",
        "assignment_reason": (
            f"Assigned to {t.driver.name if t.driver else 'Driver'} with {t.vehicle.name if t.vehicle else 'Vehicle'} "
            f"based on capacity feasibility ({t.vehicle.capacity if t.vehicle else 'N/A'}), availability, and proximity."
        ),
        "best_route_json": {
            "distance_km": t.total_distance_km or 0.0,
            "duration_min": t.total_duration_min or 0.0,
            "stops": enriched_stops,
        },
        "chosen_provider": (
            t.driver.provider.name
            if t.driver is not None and t.driver.provider is not None
            else "Default Provider"
        ),
        "chosen_vehicle": t.vehicle.name if t.vehicle else "Unknown",
        "estimated_cost": t.estimated_cost or 0.0,
        "eta_mins": t.eta_min or 0.0,
        "fuel_saved_l": t.fuel_saved_l or 0.0,
        "distance_saved_km": t.distance_saved_km or 0.0,
        "co2_saved_kg": t.co2_saved_kg or 0.0,
        "optimization_score": t.optimization_score or 0.0,
        "compatibility_score": compatibility_score,
        "factor_scores": factor_scores,
        "factor_details": factor_details,
        "reasons": reasons,
        "explanation_json": {
            "status": "Accepted",
            "decision": "Feasible",
            "type": "shared" if t.is_shared else "individual",
            "trip_code": t.trip_code,
            "optimization_score": t.optimization_score or 0.0,
        },
        "created_at": t.created_at,
    }


def _batch_to_result(b: DMFEBatch, db: SessionDep) -> dict:
    """
    Serialize one rejected A-DMFE candidate batch into result shape.
    Does NOT output a dispatched route, driver, or vehicle.
    """
    request_ids = json_loads(b.request_ids_json, [])
    reqs = (
        db.query(SimulationRequest).filter(SimulationRequest.id.in_(request_ids)).all()
        if request_ids
        else []
    )
    factor_scores = json_loads(b.factor_scores_json, {})
    factor_details = json_loads(b.factor_details_json, {})
    reasons = json_loads(b.reason_json, [])

    # The real rejection gate is always recorded as the FIRST "✗" reason by
    # the decision engine / pipeline.  Reading it directly is deterministic;
    # keyword-scanned fallback was mislabelling (e.g. capacity/weight
    # rejections landed on "rejected"/"bqs" keyword hits).
    failed_gate = "Incompatible constraints"
    for r in reasons:
        r_str = str(r)
        if r_str.startswith("✗"):
            failed_gate = r_str
            break

    return {
        "id": 1000000 + b.id,
        "raw_id": b.id,
        "is_batch": True,
        "batch_id": b.batch_code,
        "decision": "Rejected",
        "status": "Rejected",
        "is_shared": len(request_ids) > 1,
        "request_count": len(request_ids),
        "requests": [_serialize_request(r) for r in reqs],
        "provider_id": None,
        "vehicle_id": None,
        "driver": None,
        "vehicle": None,
        "current_load": None,
        "vehicle_capacity": None,
        "utilization_pct": 0.0,
        "driver_availability": "N/A — Not Dispatched",
        "assignment_reason": None,
        "chosen_provider": "None",
        "chosen_vehicle": "None",
        "estimated_cost": 0.0,
        "eta_mins": 0.0,
        "fuel_saved_l": 0.0,
        "distance_saved_km": 0.0,
        "co2_saved_kg": 0.0,
        "optimization_score": b.compatibility_score or 0.0,
        "compatibility_score": b.compatibility_score or 0.0,
        "factor_scores": factor_scores,
        "factor_details": factor_details,
        "reasons": reasons,
        "failed_gate": failed_gate,
        "best_route_json": None,
        "explanation_json": {
            "status": "Rejected",
            "decision": "Incompatible",
            "rejection_reason": failed_gate,
            "trip_code": b.batch_code,
        },
        "created_at": b.created_at,
    }


@router.get("/results")
def list_results(
    db: SessionDep,
    current_user: CurrentUser,
    limit: int = 50,
    status_filter: Optional[str] = "all",
):
    """
    Recent A-DMFE orchestration results.
    Returns dispatched trips and/or rejected batches according to status_filter ('all' | 'accepted' | 'rejected').
    """
    items = []
    if status_filter in ("all", "accepted"):
        trips = (
            db.query(Trip)
            .order_by(Trip.created_at.desc())
            .limit(limit)
            .all()
        )
        items.extend([_trip_to_result(t, db) for t in trips])

    if status_filter in ("all", "rejected"):
        batches = (
            db.query(DMFEBatch)
            .filter(DMFEBatch.status == "Rejected")
            .order_by(DMFEBatch.created_at.desc())
            .limit(limit)
            .all()
        )
        items.extend([_batch_to_result(b, db) for b in batches])

    # Sort combined items newest first
    items.sort(key=lambda x: str(x.get("created_at") or ""), reverse=True)
    return items[:limit]


@router.get("/results/{result_id}")
def get_result(result_id: str, db: SessionDep, current_user: CurrentUser):
    """One A-DMFE trip or rejected batch by ID or batch code."""
    try:
        num_id = int(result_id)
        if num_id >= 1000000:
            batch = db.query(DMFEBatch).filter(DMFEBatch.id == (num_id - 1000000)).first()
            if batch:
                return _batch_to_result(batch, db)
        else:
            trip = db.query(Trip).filter(Trip.id == num_id).first()
            if trip:
                return _trip_to_result(trip, db)
    except ValueError:
        pass

    if str(result_id).startswith("batch-"):
        try:
            b_id = int(result_id.split("-")[1])
            batch = db.query(DMFEBatch).filter(DMFEBatch.id == b_id).first()
            if batch:
                return _batch_to_result(batch, db)
        except ValueError:
            pass

    trip = db.query(Trip).filter(Trip.trip_code == result_id).first()
    if trip:
        return _trip_to_result(trip, db)

    batch = db.query(DMFEBatch).filter(DMFEBatch.batch_code == result_id).first()
    if batch:
        return _batch_to_result(batch, db)

    raise HTTPException(404, "Result not found")


@router.post("/simulate")
def simulate_requests(db: SessionDep, current_user: CurrentUser, count: int = 10):
    from app.services.mock_adapters import generate_simulation_requests
    requests = generate_simulation_requests(count, db)
    return {"message": f"{len(requests)} simulation requests created"}


@router.get("/requests")
def list_requests(db: SessionDep, current_user: CurrentUser, limit: int = 50):
    reqs = (
        db.query(SimulationRequest)
        .order_by(SimulationRequest.created_at.desc())
        .limit(limit)
        .all()
    )
    return [
        {
            "id": r.id,
            "request_type": r.request_type,
            "provider_id": r.provider_id,
            "pickup_lat": r.pickup_lat,
            "pickup_lng": r.pickup_lng,
            "drop_lat": r.drop_lat,
            "drop_lng": r.drop_lng,
            "pickup_address": r.pickup_address or "",
            "drop_address": r.drop_address or "",
            "demand": r.demand or 1,
            "weight_kg": r.weight_kg or 0.0,
            "priority": r.priority or "Medium",
            "vehicle_type": r.vehicle_type or "Auto",
            "estimated_distance_km": r.estimated_distance_km or 0.0,
            "status": r.status or "Pending",
            "created_at": r.created_at,
        }
        for r in reqs
    ]
